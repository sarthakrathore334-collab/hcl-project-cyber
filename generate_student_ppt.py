import os
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE

def build_student_presentation():
    prs = Presentation()
    # 16:9 Widescreen standard
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6]

    # Warm, clean, aesthetic color palette
    COLOR_BG = RGBColor(247, 245, 240)          # #F7F5F0 (Warm off-white background)
    COLOR_CARD = RGBColor(255, 255, 255)        # #FFFFFF (Clean card surface)
    COLOR_CARD_ALT = RGBColor(240, 237, 230)    # #F0EDE6 (Accent card background)
    COLOR_BORDER = RGBColor(227, 222, 213)      # #E3DED5 (Subtle border)
    COLOR_TEXT = RGBColor(28, 25, 23)           # #1C1917 (Primary text)
    COLOR_TEXT_MUTED = RGBColor(120, 113, 108)  # #78716C (Secondary text)
    COLOR_ACCENT = RGBColor(146, 64, 14)        # #92400E (Warm terracotta amber)
    COLOR_DANGER = RGBColor(153, 27, 27)        # #991B1B (High threat red)
    COLOR_SAFE = RGBColor(22, 101, 52)          # #166534 (Safe green)
    COLOR_BLUE = RGBColor(30, 64, 175)          # #1E40AF (Info blue)

    FONT_TITLE = "Georgia"
    FONT_BODY = "Arial"

    def set_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = COLOR_BG
        bg.line.color.rgb = COLOR_BG
        return bg

    def add_header(slide, tag, title, subtitle):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.45), Inches(11.733), Inches(1.2))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p_tag = tf.paragraphs[0]
        p_tag.text = tag.upper()
        p_tag.font.size = Pt(11)
        p_tag.font.bold = True
        p_tag.font.name = FONT_BODY
        p_tag.font.color.rgb = COLOR_ACCENT

        p_title = tf.add_paragraph()
        p_title.text = title
        p_title.font.size = Pt(26)
        p_title.font.bold = True
        p_title.font.name = FONT_TITLE
        p_title.font.color.rgb = COLOR_TEXT
        p_title.space_before = Pt(3)

        p_sub = tf.add_paragraph()
        p_sub.text = subtitle
        p_sub.font.size = Pt(13)
        p_sub.font.name = FONT_BODY
        p_sub.font.color.rgb = COLOR_TEXT_MUTED
        p_sub.space_before = Pt(3)

    def add_card(slide, left, top, width, height, bg_color=COLOR_CARD, border_color=COLOR_BORDER):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = bg_color
        card.line.color.rgb = border_color
        card.line.width = Pt(1)
        return card

    # =========================================================================
    # SLIDE 1: Title Slide (Student & Mentor Details)
    # =========================================================================
    s1 = prs.slides.add_slide(blank_layout)
    set_bg(s1)

    add_card(s1, 1.0, 0.9, 11.333, 5.7, bg_color=COLOR_CARD)

    tb = s1.shapes.add_textbox(Inches(1.5), Inches(1.3), Inches(10.333), Inches(3.0))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "ACADEMIC CYBERSECURITY PROJECT PRESENTATION"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT

    p = tf.add_paragraph()
    p.text = "XSS Guard"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT
    p.space_before = Pt(6)

    p = tf.add_paragraph()
    p.text = "Automated Cross-Site Scripting Detection, Real-Time Reflection Testing & Security Auditing"
    p.font.size = Pt(17)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    # Student & Mentor Info Box
    info_card = add_card(s1, 1.5, 4.3, 10.333, 1.8, bg_color=COLOR_CARD_ALT)
    tb_info = s1.shapes.add_textbox(Inches(1.75), Inches(4.45), Inches(9.833), Inches(1.5))
    tf_info = tb_info.text_frame
    tf_info.word_wrap = True

    p = tf_info.paragraphs[0]
    p.text = "PROJECT METADATA & PRESENTATION DETAILS"
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT

    p = tf_info.add_paragraph()
    p.text = "• Presented By:  [ Student Name / Roll No / Group ID ]\n• Project Mentor: [ Mentor / Faculty Guide / Professor Name ]\n• Department / Subject: Computer Science & Engineering • Web Application Security"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT
    p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 2: Simple Introduction - What is XSS?
    # =========================================================================
    s2 = prs.slides.add_slide(blank_layout)
    set_bg(s2)
    add_header(s2, "Introduction", "What is Cross-Site Scripting (XSS)?", "Understanding the vulnerability in simple, practical terms.")

    c1 = add_card(s2, 0.8, 1.85, 3.65, 5.0)
    tb = s2.shapes.add_textbox(Inches(1.05), Inches(2.05), Inches(3.15), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The Concept"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "Cross-Site Scripting (XSS) is a code injection flaw where an attacker injects malicious client-side scripts into trusted web pages.\n\nBecause the browser trusts the web application, it executes the attacker's script without questioning its validity."
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(10)

    c2 = add_card(s2, 4.84, 1.85, 3.65, 5.0)
    tb = s2.shapes.add_textbox(Inches(5.09), Inches(2.05), Inches(3.15), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "The Everyday Analogy"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "Think of a trusted courier service:\n\nIf someone sneaks a forged letter into an official envelope, the recipient opens and trusts it because the envelope looks 100% genuine.\n\nIn XSS, untrusted JavaScript rides inside a genuine website's HTML!"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(10)

    c3 = add_card(s2, 8.88, 1.85, 3.65, 5.0)
    tb = s2.shapes.add_textbox(Inches(9.13), Inches(2.05), Inches(3.15), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Why It Matters (OWASP)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "• Ranked in the OWASP Top 10 web security risks.\n• Can hijack user sessions (stealing login cookies).\n• Can capture keystrokes & passwords.\n• Can silently redirect users to phishing sites."
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 3: The 3 Main Types of XSS
    # =========================================================================
    s3 = prs.slides.add_slide(blank_layout)
    set_bg(s3)
    add_header(s3, "Taxonomy", "The Three Types of XSS Attacks", "How attackers deliver malicious scripts to victims.")

    types = [
        ("1. Reflected XSS", "Non-Persistent Attack", "The malicious script is reflected off the web server immediately in a search query or error message.\n\nExample: An attacker sends a link with ?q=<script>alert(1)</script>. When clicked, the server echoes the script directly back to the victim's browser."),
        ("2. Stored XSS", "Persistent Threat", "The malicious script is permanently stored in the database (e.g. comments, forum posts, user profiles).\n\nExample: An attacker posts a malicious payload in a review. Every subsequent user who views the page automatically runs the payload!"),
        ("3. DOM-based XSS", "Client-Side Execution", "The attack occurs entirely within the victim's browser by manipulating the Document Object Model (DOM).\n\nExample: JavaScript reads an unsanitized fragment from location.hash and directly inserts it into the DOM via innerHTML.")
    ]

    for i, (title, tag, body) in enumerate(types):
        left = 0.8 + (i * 3.98)
        add_card(s3, left, 1.85, 3.75, 5.0)
        tb = s3.shapes.add_textbox(Inches(left + 0.25), Inches(2.05), Inches(3.25), Inches(4.5))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.name = FONT_TITLE
        p.font.color.rgb = COLOR_TEXT

        p = tf.add_paragraph()
        p.text = tag.upper()
        p.font.size = Pt(10)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT
        p.space_before = Pt(3)

        p = tf.add_paragraph()
        p.text = body
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(12)

    # =========================================================================
    # SLIDE 4: Project Objectives & Student Motivation
    # =========================================================================
    s4 = prs.slides.add_slide(blank_layout)
    set_bg(s4)
    add_header(s4, "Project Scope", "Project Goals & What We Built", "Our response: An integrated, practical XSS testing and incident response toolkit.")

    goals = [
        ("01", "Dual-Stage Threat Detection", "Instant in-browser heuristic checking (<10ms) paired with authoritative backend signature analysis."),
        ("02", "Active Reflection Scanner", "Safe, non-destructive probe tester that inspects whether web endpoints echo user parameters unsanitized."),
        ("03", "SOC Operations Telemetry", "Interactive dashboard with real-time counters, search query filtering, and single-click CSV log downloads."),
        ("04", "Formal Audit Compliance", "Automated executive risk score (Critical, Moderate, Low) with PDF-ready print stylesheets and reports.")
    ]

    for i, (num, heading, text) in enumerate(goals):
        left = 0.8 + (i * 2.98)
        add_card(s4, left, 1.85, 2.78, 5.0)
        tb = s4.shapes.add_textbox(Inches(left + 0.2), Inches(2.1), Inches(2.38), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = num
        p.font.size = Pt(28)
        p.font.bold = True
        p.font.name = FONT_TITLE
        p.font.color.rgb = COLOR_ACCENT

        p = tf.add_paragraph()
        p.text = heading
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.name = FONT_TITLE
        p.font.color.rgb = COLOR_TEXT
        p.space_before = Pt(6)

        p = tf.add_paragraph()
        p.text = text
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 5: Architecture & Data Flow
    # =========================================================================
    s5 = prs.slides.add_slide(blank_layout)
    set_bg(s5)
    add_header(s5, "Architecture", "System Architecture & Data Pipeline", "How data flows through the application from input to telemetry.")

    flow_steps = [
        ("Stage 1: User Input", "User submits raw payload or target endpoint for testing."),
        ("Stage 2: Heuristic Analysis", "detector.py checks for High/Medium signatures (script, onerror, cookie)."),
        ("Stage 3: Reflection Probe", "scanner.py injects controlled token (XSS_TEST_123<test>) into target."),
        ("Stage 4: Logging & Storage", "All confirmed suspicious incidents are appended to logs.txt."),
        ("Stage 5: SOC & Reporting", "Dashboard aggregates metrics, enables search, CSV export & audit reports.")
    ]

    for i, (title, desc) in enumerate(flow_steps):
        left = 0.8 + (i * 2.38)
        add_card(s5, left, 1.85, 2.2, 5.0)
        tb = s5.shapes.add_textbox(Inches(left + 0.15), Inches(2.1), Inches(1.9), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"STEP {i+1}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT

        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(14)
        p.font.bold = True
        p.font.name = FONT_TITLE
        p.font.color.rgb = COLOR_TEXT
        p.space_before = Pt(6)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 6: Feature 1 - Real-Time XSS Detector
    # =========================================================================
    s6 = prs.slides.add_slide(blank_layout)
    set_bg(s6)
    add_header(s6, "Core Feature 1", "Real-Time XSS Detector Module", "Instant client-side heuristic feedback combined with server-side validation.")

    add_card(s6, 0.8, 1.85, 5.75, 5.0)
    tb = s6.shapes.add_textbox(Inches(1.05), Inches(2.1), Inches(5.25), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Detection Rules & Severity Matrix"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "• High Severity Threats:\n  - <script : Direct HTML script tag execution\n  - javascript: : Inline pseudo-protocol links\n  - document.cookie : Session hijacking / cookie exfiltration\n\n• Medium Severity Threats:\n  - onerror=, onload=, onclick= : Inline DOM event triggers\n  - <iframe, <svg : Obfuscated tag embedding containers\n\n• Low Severity / Safe:\n  - Normal text strings or benign inputs without executable triggers."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    add_card(s6, 6.78, 1.85, 5.75, 5.0)
    tb = s6.shapes.add_textbox(Inches(7.03), Inches(2.1), Inches(5.25), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Interactive Features for Mentors"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "• 1-Click Sample Payload Chips:\n  Built-in sample vector pills allow testing with a single click during live presentations without manual typing!\n\n• Live Character & Pattern Counter:\n  Shows instant character count and live threat status badge.\n\n• Server Logging:\n  Confirmed threats are automatically logged into logs.txt with timestamp and pattern metadata."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 7: Feature 2 - Active Reflection Scanner & Lab Testing
    # =========================================================================
    s7 = prs.slides.add_slide(blank_layout)
    set_bg(s7)
    add_header(s7, "Core Feature 2", "Active Reflection Scanner & Test Lab", "Validating whether query parameters reflect unescaped user input.")

    add_card(s7, 0.8, 1.85, 5.75, 5.0)
    tb = s7.shapes.add_textbox(Inches(1.05), Inches(2.1), Inches(5.25), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "How the Scanner Works"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "1. Non-Destructive Probe Injection:\n   Sends safe probe token: XSS_TEST_123<test>\n\n2. Reflection Context Inspection:\n   Checks if < and > are reflected unescaped.\n\n3. Classification:\n   - High: Echoed raw inside HTML body.\n   - Medium: Echoed inside HTML attribute.\n   - Safe: Encoded as &lt;test&gt; or stripped."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    add_card(s7, 6.78, 1.85, 5.75, 5.0)
    tb = s7.shapes.add_textbox(Inches(7.03), Inches(2.1), Inches(5.25), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Our Test Lab: Vulnerable vs Safe"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "We built test_app.py to prove both cases:\n\n• Vulnerable Endpoint (/search):\n  <p>You searched for: {query}</p>\n  [!] Scanner flags as HIGH: raw reflection!\n\n• Protected Endpoint (/safe):\n  from markupsafe import escape\n  <p>You searched for: {escape(query)}</p>\n  [✓] Scanner confirms SAFE: characters sanitized!"
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 8: Feature 3 - SOC Telemetry Dashboard
    # =========================================================================
    s8 = prs.slides.add_slide(blank_layout)
    set_bg(s8)
    add_header(s8, "Core Feature 3", "Security Operations (SOC) Dashboard", "Transforming raw logs into actionable incident telemetry.")

    soc_cards = [
        ("Animated KPI Counters", "Visualizes Total, High, Medium, and Low detections with animated counter transitions upon page load."),
        ("Proportional Risk Meter", "Calculates percentage distribution across severity tiers to show current threat exposure at a glance."),
        ("Instant Real-Time Search", "Type-to-filter search bar filtering logs instantly by user input keyword or matched pattern name."),
        ("One-Click CSV Export", "Directly generates and downloads xss-detections.csv for external SIEM integration and mentor grading.")
    ]

    for i, (title, desc) in enumerate(soc_cards):
        col = i % 2
        row = i // 2
        left = 0.8 + (col * 5.95)
        top = 1.85 + (row * 2.5)
        add_card(s8, left, top, 5.75, 2.3)

        tb = s8.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(5.25), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.name = FONT_TITLE
        p.font.color.rgb = COLOR_TEXT

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 9: Feature 4 - Learn Center & Formal Audit Report
    # =========================================================================
    s9 = prs.slides.add_slide(blank_layout)
    set_bg(s9)
    add_header(s9, "Core Feature 4", "Educational Center & Formal Audit Report", "Bridging theoretical knowledge with practical vulnerability compliance.")

    add_card(s9, 0.8, 1.85, 5.75, 5.0)
    tb = s9.shapes.add_textbox(Inches(1.05), Inches(2.1), Inches(5.25), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Learn Center (/learn)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "• Student & Developer Reference Guide:\n  - Detailed breakdown of Reflected, Stored, and DOM-based XSS.\n  - Real-world case studies (Samy Worm, Yahoo Mail, Twitter StalkDaily).\n  - Concrete remediation cheat sheets.\n\n• Interactive in-page table of contents for rapid presentation navigation."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    add_card(s9, 6.78, 1.85, 5.75, 5.0)
    tb = s9.shapes.add_textbox(Inches(7.03), Inches(2.1), Inches(5.25), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Security Audit Report (/report)"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "• Automated Executive Risk Verdict:\n  Classifies security posture as Critical, Moderate, or Low based on findings.\n\n• Chronological Audit Table:\n  Numbered list of exact attack payloads, patterns, and timestamps.\n\n• Dual Export Capability:\n  - Standalone offline HTML file download.\n  - Print / Save as PDF with dedicated print CSS."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 10: Live Demonstration Flow (For Mentor Presentation)
    # =========================================================================
    s10 = prs.slides.add_slide(blank_layout)
    set_bg(s10)
    add_header(s10, "Classroom Demo", "Recommended Live Presentation Flow", "A 4-step sequence to demonstrate the project smoothly to mentors.")

    demo_steps = [
        ("Step 1: Test Detection", "Open / in browser.\nClick the <script> chip or type a payload.\nShow instant live check and server detection log."),
        ("Step 2: Run Scanner", "Open /scan.\nClick 'Use local test' (points to test_app.py).\nShow probe animation and reflection context snippet."),
        ("Step 3: Review Dashboard", "Open /dashboard.\nShow animated counters, filter by 'High', use search bar, and download CSV spreadsheet."),
        ("Step 4: Show Report", "Open /report.\nShow executive risk rating (Critical/Moderate), then click 'Print / Save as PDF'.")
    ]

    for i, (title, desc) in enumerate(demo_steps):
        left = 0.8 + (i * 2.98)
        add_card(s10, left, 1.85, 2.78, 5.0)
        tb = s10.shapes.add_textbox(Inches(left + 0.2), Inches(2.1), Inches(2.38), Inches(4.4))
        tf = tb.text_frame
        tf.word_wrap = True

        p = tf.paragraphs[0]
        p.text = f"DEMO {i+1}"
        p.font.size = Pt(11)
        p.font.bold = True
        p.font.color.rgb = COLOR_ACCENT

        p = tf.add_paragraph()
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.name = FONT_TITLE
        p.font.color.rgb = COLOR_TEXT
        p.space_before = Pt(6)

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(10)

    # =========================================================================
    # SLIDE 11: Mitigation & Prevention (The Solution)
    # =========================================================================
    s11 = prs.slides.add_slide(blank_layout)
    set_bg(s11)
    add_header(s11, "Defense In Depth", "How Developers Prevent XSS", "Industry standard defensive controls recommended by OWASP.")

    fixes = [
        ("Context-Aware Output Encoding", "Convert dangerous characters into safe HTML entities before rendering in browser (&lt;, &gt;, &quot;, &#x27;)."),
        ("Content Security Policy (CSP)", "Configure HTTP headers restricting executable script domains and disabling inline execution (script-src 'self')."),
        ("Cookie Hardening (HttpOnly)", "Mark session cookies as HttpOnly so JavaScript cannot access them even if an XSS flaw exists."),
        ("Auto-Escaping Frameworks", "Use modern template engines (Jinja2, React, Angular) that automatically escape variables by default.")
    ]

    for i, (title, desc) in enumerate(fixes):
        col = i % 2
        row = i // 2
        left = 0.8 + (col * 5.95)
        top = 1.85 + (row * 2.5)
        add_card(s11, left, top, 5.75, 2.3)

        tb = s11.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.2), Inches(5.25), Inches(1.8))
        tf = tb.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title
        p.font.size = Pt(16)
        p.font.bold = True
        p.font.name = FONT_TITLE
        p.font.color.rgb = COLOR_TEXT

        p = tf.add_paragraph()
        p.text = desc
        p.font.size = Pt(12)
        p.font.color.rgb = COLOR_TEXT_MUTED
        p.space_before = Pt(6)

    # =========================================================================
    # SLIDE 12: Learning Outcomes & Future Scope
    # =========================================================================
    s12 = prs.slides.add_slide(blank_layout)
    set_bg(s12)
    add_header(s12, "Summary", "Key Learnings & Future Enhancements", "Academic takeaways and future technical directions.")

    add_card(s12, 0.8, 1.85, 5.75, 5.0)
    tb = s12.shapes.add_textbox(Inches(1.05), Inches(2.1), Inches(5.25), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "What We Learned as Students"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "• Deep understanding of OWASP Top 10 vulnerabilities.\n• Hands-on socket and HTTP probe analysis in Python.\n• Full-stack integration connecting Python Flask backend with responsive, minimalist UI.\n• Practical understanding of why sanitization and context-aware encoding are essential.\n• Creating audit reports for compliance readiness."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    add_card(s12, 6.78, 1.85, 5.75, 5.0)
    tb = s12.shapes.add_textbox(Inches(7.03), Inches(2.1), Inches(5.25), Inches(4.5))
    tf = tb.text_frame
    tf.word_wrap = True
    p = tf.paragraphs[0]
    p.text = "Future Technical Scope"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT

    p = tf.add_paragraph()
    p.text = "• Machine Learning Model: Train NLP classifiers to detect polyglot and obfuscated XSS payloads.\n• CI/CD Pipeline Scanning: Automate scanning on git push via GitHub Actions.\n• Web Application Firewall (WAF) Proxy: Intercept and drop malicious payloads before they hit backend servers.\n• Instant Webhook Alerts: Route alerts to Slack or Discord."
    p.font.size = Pt(12)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    # =========================================================================
    # SLIDE 13: Conclusion / Q&A Slide
    # =========================================================================
    s13 = prs.slides.add_slide(blank_layout)
    set_bg(s13)

    add_card(s13, 1.5, 1.2, 10.333, 5.1, bg_color=COLOR_CARD)
    tb = s13.shapes.add_textbox(Inches(2.0), Inches(1.8), Inches(9.333), Inches(3.8))
    tf = tb.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "CONCLUSION & PROJECT CLOSEOUT"
    p.font.size = Pt(12)
    p.font.bold = True
    p.font.color.rgb = COLOR_ACCENT

    p = tf.add_paragraph()
    p.text = "Thank You!"
    p.font.size = Pt(46)
    p.font.bold = True
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "Questions, Feedback & Mentor Discussion"
    p.font.size = Pt(20)
    p.font.name = FONT_TITLE
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(8)

    p = tf.add_paragraph()
    p.text = "XSS Guard — Empowering developers and analysts with transparent, proactive web application security.\n\nGitHub / Repository: [ Project Workspace ]  •  Live Demo: http://127.0.0.1:5000"
    p.font.size = Pt(13)
    p.font.color.rgb = COLOR_TEXT_MUTED
    p.space_before = Pt(16)

    # Save presentation
    filename = "XSS_Guard_Student_Presentation.pptx"
    prs.save(filename)
    print(f"Student presentation successfully created: {filename}")

if __name__ == "__main__":
    build_student_presentation()
